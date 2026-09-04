import hashlib
import io
import ipaddress
import os
import re
import socket
import subprocess
import tempfile
import zipfile
from urllib.parse import urlparse

import requests
import streamlit as st
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter


# ============================================================
# APP CONFIG
# ============================================================

st.set_page_config(
    page_title="DocumentSentinel",
    page_icon="🛡️",
    layout="wide",
)

st.title("🛡️ DocumentSentinel")
st.caption(
    "Multi-format document forensics, threat indicators, and RAG investigation"
)

SUPPORTED = [".pdf", ".docx", ".xlsx"]


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("⚙️ Agent Settings")

    agent_type = st.selectbox(
        "Choose Agent",
        [
            "🔍 Forensic Analyzer",
            "🤖 RAG Investigator",
            "🦠 Security Analyzer",
            "🧩 Full Investigation",
        ],
    )

    st.divider()

    st.subheader("API Keys")
    openai_key = st.text_input(
        "OpenAI API Key",
        value=os.getenv("OPENAI_API_KEY", ""),
        type="password",
    )
    virustotal_key = st.text_input(
        "VirusTotal API Key (optional)",
        value="",
        type="password",
    )

    st.divider()

    st.subheader("Analysis")
    do_metadata = st.checkbox("Metadata", True)
    do_hash = st.checkbox("SHA-256 Hash", True)
    do_structure = st.checkbox("Suspicious Structures", True)
    do_iocs = st.checkbox("URLs / IOCs", True)
    do_clamav = st.checkbox("ClamAV Scan", True)
    do_vt = st.checkbox("VirusTotal Hash Check", False)

    st.divider()
    st.caption(
        "Safety note: the application analyzes files without opening them "
        "in Microsoft Office or executing embedded content."
    )


# ============================================================
# HELPERS
# ============================================================

def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def safe_filename(name: str) -> str:
    name = os.path.basename(name or "document")
    return re.sub(r"[^A-Za-z0-9._-]", "_", name)


def validate_pdf_url(url: str) -> tuple[bool, str]:
    """Basic SSRF protection for the URL downloader."""
    try:
        parsed = urlparse(url)

        if parsed.scheme not in {"http", "https"}:
            return False, "Only HTTP and HTTPS URLs are allowed."

        if not parsed.hostname:
            return False, "The URL has no valid hostname."

        host = parsed.hostname

        # Reject obvious local hostnames.
        if host.lower() in {
            "localhost",
            "localhost.localdomain",
            "metadata.google.internal",
        } or host.endswith(".local"):
            return False, "Local/internal hostnames are not allowed."

        # Reject literal private/reserved IPs.
        try:
            ip = ipaddress.ip_address(host)
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_reserved
                or ip.is_multicast
            ):
                return False, "Private or reserved IP addresses are not allowed."
        except ValueError:
            # Domain name; DNS is resolved below.
            pass

        # Resolve the hostname and reject private/reserved addresses.
        for result in socket.getaddrinfo(host, None):
            resolved = result[4][0]
            try:
                ip = ipaddress.ip_address(resolved)
                if (
                    ip.is_private
                    or ip.is_loopback
                    or ip.is_link_local
                    or ip.is_reserved
                    or ip.is_multicast
                ):
                    return False, "The URL resolves to a private/reserved address."
            except ValueError:
                continue

        return True, ""
    except Exception as exc:
        return False, f"Invalid URL: {exc}"


def download_document(url: str) -> bytes:
    ok, reason = validate_pdf_url(url)
    if not ok:
        raise ValueError(reason)

    response = requests.get(
        url,
        timeout=20,
        headers={"User-Agent": "DocumentSentinel/1.0"},
        allow_redirects=False,
    )
    response.raise_for_status()

    content_type = response.headers.get("content-type", "").lower()
    if (
        "application/pdf" not in content_type
        and not url.lower().split("?")[0].endswith(".pdf")
    ):
        raise ValueError("The URL does not appear to point directly to a PDF.")

    if len(response.content) > 25 * 1024 * 1024:
        raise ValueError("The downloaded file is larger than 25 MB.")

    return response.content


# ============================================================
# CONTENT EXTRACTION
# ============================================================

def extract_pdf(pdf_bytes: bytes):
    documents = []
    metadata = {}
    path = None

    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(pdf_bytes)
            path = tmp.name

        reader = PdfReader(path)

        metadata["File type"] = "PDF"
        metadata["Pages"] = len(reader.pages)
        metadata["Encrypted"] = reader.is_encrypted

        if reader.metadata:
            for key, value in reader.metadata.items():
                metadata[str(key).lstrip("/")] = str(value)

        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": "PDF", "page": index + 1},
                    )
                )

    finally:
        if path and os.path.exists(path):
            os.unlink(path)

    return documents, metadata


def extract_docx(docx_bytes: bytes):
    metadata = {"File type": "Word (.docx)"}
    documents = []

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        names = z.namelist()

        metadata["ZIP members"] = len(names)
        metadata["Embedded files"] = sum(
            n.startswith("word/embeddings/") for n in names
        )
        metadata["Images"] = sum(n.startswith("word/media/") for n in names)
        metadata["VBA project"] = any(
            n.endswith("vbaProject.bin") for n in names
        )

    doc = DocxDocument(io.BytesIO(docx_bytes))

    props = doc.core_properties
    for label, value in {
        "Title": props.title,
        "Subject": props.subject,
        "Author": props.author,
        "Last modified by": props.last_modified_by,
        "Created": props.created,
        "Modified": props.modified,
        "Keywords": props.keywords,
        "Comments": props.comments,
    }.items():
        if value:
            metadata[label] = str(value)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    for i in range(0, len(paragraphs), 30):
        block = "\n".join(paragraphs[i:i + 30])
        if block.strip():
            documents.append(
                Document(
                    page_content=block,
                    metadata={"source": "DOCX", "section": i // 30 + 1},
                )
            )

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text for cell in row.cells))
        content = "\n".join(rows)
        if content.strip():
            documents.append(
                Document(
                    page_content=content,
                    metadata={"source": "DOCX table", "table": table_index},
                )
            )

    return documents, metadata


def extract_xlsx(xlsx_bytes: bytes):
    metadata = {"File type": "Excel (.xlsx)"}
    documents = []

    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as z:
        names = z.namelist()

        metadata["ZIP members"] = len(names)
        metadata["Embedded files"] = sum(
            n.startswith("xl/embeddings/") for n in names
        )
        metadata["VBA project"] = any(
            n.endswith("vbaProject.bin") for n in names
        )
        metadata["External links"] = sum(
            n.startswith("xl/externalLinks/") for n in names
        )

    workbook = load_workbook(
        io.BytesIO(xlsx_bytes),
        read_only=True,
        data_only=False,
    )

    metadata["Sheets"] = len(workbook.sheetnames)
    metadata["Sheet names"] = ", ".join(workbook.sheetnames)

    props = workbook.properties
    for label, value in {
        "Title": props.title,
        "Subject": props.subject,
        "Creator": props.creator,
        "Last modified by": props.lastModifiedBy,
        "Created": props.created,
        "Modified": props.modified,
        "Keywords": props.keywords,
    }.items():
        if value:
            metadata[label] = str(value)

    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(" | ".join(values))

            if len(lines) >= 250:
                break

        content = "\n".join(lines)

        if content.strip():
            documents.append(
                Document(
                    page_content=content,
                    metadata={"source": "XLSX", "sheet": sheet.title},
                )
            )

    workbook.close()
    return documents, metadata


def extract_content(file_name: str, data: bytes):
    ext = os.path.splitext(file_name)[1].lower()

    if ext == ".pdf":
        return extract_pdf(data)

    if ext == ".docx":
        return extract_docx(data)

    if ext == ".xlsx":
        return extract_xlsx(data)

    raise ValueError(
        "Unsupported file type. Use PDF, DOCX, or XLSX."
    )


# ============================================================
# FORENSIC ANALYSIS
# ============================================================

def analyze_pdf_structure(data: bytes):
    text = data.decode("latin-1", errors="ignore")

    patterns = {
        "/JavaScript": "JavaScript reference",
        "/JS": "JavaScript object/reference",
        "/OpenAction": "Automatic open action",
        "/AA": "Additional automatic action",
        "/Launch": "Launch action",
        "/EmbeddedFile": "Embedded file",
        "/RichMedia": "Rich media object",
        "/AcroForm": "Interactive form",
        "/SubmitForm": "Form submission action",
        "/GoToR": "Remote document action",
        "/URI": "URI action",
    }

    findings = []
    for pattern, description in patterns.items():
        count = text.count(pattern)
        if count:
            findings.append(
                {
                    "indicator": pattern,
                    "description": description,
                    "count": count,
                }
            )

    return findings


def analyze_office_structure(data: bytes, extension: str):
    findings = []

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            lower_names = [n.lower() for n in names]

            if any(n.endswith("vbaproject.bin") for n in lower_names):
                findings.append(
                    {
                        "indicator": "VBA project",
                        "description": "A VBA project is present.",
                        "count": 1,
                    }
                )

            embedded_prefix = (
                "word/embeddings/" if extension == ".docx"
                else "xl/embeddings/"
            )
            embedded = [
                n for n in names if n.lower().startswith(embedded_prefix)
            ]
            if embedded:
                findings.append(
                    {
                        "indicator": "Embedded objects",
                        "description": "Embedded object(s) are present.",
                        "count": len(embedded),
                    }
                )

            external_prefix = "xl/externallinks/"
            external = [
                n for n in lower_names if n.startswith(external_prefix)
            ]
            if external:
                findings.append(
                    {
                        "indicator": "External links",
                        "description": "External-link package entries are present.",
                        "count": len(external),
                    }
                )

            suspicious_names = [
                n for n in lower_names
                if any(
                    token in n
                    for token in (
                        "oleobject",
                        "activex",
                        "external",
                    )
                )
            ]

            if suspicious_names:
                findings.append(
                    {
                        "indicator": "Potentially interesting OOXML objects",
                        "description": (
                            "The document package contains object/link-related "
                            "components that deserve review."
                        ),
                        "count": len(suspicious_names),
                    }
                )

    except zipfile.BadZipFile:
        findings.append(
            {
                "indicator": "Invalid ZIP package",
                "description": "The Office document is not a valid ZIP-based package.",
                "count": 1,
            }
        )

    return findings


def extract_iocs(data: bytes, documents):
    raw = data.decode("latin-1", errors="ignore")

    # Also search extracted human-readable text.
    extracted_text = "\n".join(
        doc.page_content for doc in documents
    )

    combined = raw + "\n" + extracted_text

    url_pattern = r"https?://[^\s<>'\"\\]+"
    ipv4_pattern = (
        r"\b(?:"
        r"(?:25[0-5]|2[0-4]\d|1?\d?\d)\."
        r"){3}"
        r"(?:25[0-5]|2[0-4]\d|1?\d?\d)\b"
    )

    urls = sorted(set(re.findall(url_pattern, combined)))
    ips = sorted(set(re.findall(ipv4_pattern, combined)))
    sha256 = sorted(set(re.findall(r"\b[a-fA-F0-9]{64}\b", combined)))
    md5 = sorted(set(re.findall(r"\b[a-fA-F0-9]{32}\b", combined)))

    return {
        "URLs": urls,
        "IPv4": ips,
        "SHA-256": sha256,
        "MD5": md5,
    }


# ============================================================
# ANTIVIRUS
# ============================================================

def clamav_scan(data: bytes, file_name: str):
    path = None

    try:
        suffix = os.path.splitext(file_name)[1] or ".bin"

        with tempfile.NamedTemporaryFile(
            suffix=suffix,
            delete=False,
        ) as tmp:
            tmp.write(data)
            path = tmp.name

        result = subprocess.run(
            ["clamscan", "--no-summary", path],
            capture_output=True,
            text=True,
            timeout=60,
        )

        output = (result.stdout + "\n" + result.stderr).strip()

        if result.returncode == 0:
            status = "Clean"
        elif result.returncode == 1:
            status = "Threat Detected"
        else:
            status = "Scan Error"

        return {"status": status, "details": output}

    except FileNotFoundError:
        return {
            "status": "ClamAV Not Installed",
            "details": (
                "Install ClamAV locally. On Streamlit Community Cloud, "
                "packages.txt can install the Linux dependency."
            ),
        }

    except Exception as exc:
        return {
            "status": "Scan Error",
            "details": str(exc),
        }

    finally:
        if path and os.path.exists(path):
            os.unlink(path)


def virustotal_hash_check(file_hash: str, api_key: str):
    if not api_key:
        return {
            "status": "Skipped",
            "details": "No VirusTotal API key supplied.",
        }

    try:
        response = requests.get(
            f"https://www.virustotal.com/api/v3/files/{file_hash}",
            headers={"x-apikey": api_key},
            timeout=20,
        )

        if response.status_code == 404:
            return {
                "status": "Unknown",
                "details": "The hash was not found in VirusTotal.",
            }

        response.raise_for_status()

        data = response.json()
        stats = (
            data.get("data", {})
            .get("attributes", {})
            .get("last_analysis_stats", {})
        )

        malicious = int(stats.get("malicious", 0))
        suspicious = int(stats.get("suspicious", 0))

        return {
            "status": (
                "Potential Threat" if malicious
                else "Suspicious" if suspicious
                else "No malicious detections"
            ),
            "malicious": malicious,
            "suspicious": suspicious,
            "details": stats,
        }

    except Exception as exc:
        return {
            "status": "VirusTotal Error",
            "details": str(exc),
        }


# ============================================================
# RISK SCORE
# ============================================================

def risk_assessment(structure, clamav_result, vt_result):
    score = min(len(structure) * 10, 40)

    if clamav_result and clamav_result.get("status") == "Threat Detected":
        score += 60

    if vt_result:
        score += min(int(vt_result.get("malicious", 0)) * 10, 50)
        score += min(int(vt_result.get("suspicious", 0)) * 5, 20)

    score = min(score, 100)

    if score >= 70:
        return "🔴 HIGH", score
    if score >= 30:
        return "🟠 MEDIUM", score
    return "🟢 LOW", score


# ============================================================
# RAG
# ============================================================

def build_vectorstore(documents, api_key, chunk_size=1000, overlap=150):
    if not api_key:
        raise ValueError("OpenAI API key is required for RAG.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    chunks = splitter.split_documents(documents)

    embeddings = OpenAIEmbeddings(
        model="text-embedding-3-small",
        api_key=api_key,
    )

    return FAISS.from_documents(chunks, embeddings)


def answer_with_rag(vectorstore, question, api_key, model_name):
    retriever = vectorstore.as_retriever(
        search_kwargs={"k": 4}
    )

    prompt = ChatPromptTemplate.from_messages(
        [
            (
                "system",
                """You are a document investigation assistant.

Answer ONLY from the retrieved document context.
Do not invent facts.
If the answer is not supported by the context, say:
"I couldn't find that in the document."

Mention the relevant source location when available
(for example, PDF page, Word section, or Excel sheet).

Retrieved context:
{context}""",
            ),
            ("human", "{question}"),
        ]
    )

    docs = retriever.invoke(question)

    context = "\n\n".join(
        f"[{doc.metadata}] {doc.page_content}"
        for doc in docs
    )

    llm = ChatOpenAI(
        model=model_name,
        temperature=0,
        api_key=api_key,
    )

    chain = prompt | llm | StrOutputParser()

    answer = chain.invoke(
        {
            "context": context,
            "question": question,
        }
    )

    return answer, docs


# ============================================================
# DOCUMENT INPUT
# ============================================================

st.header("📥 Document Input")

source_mode = st.radio(
    "Choose input",
    ["Upload document", "Online PDF URL"],
    horizontal=True,
)

file_name = None
file_bytes = None

if source_mode == "Upload document":
    uploaded = st.file_uploader(
        "Upload PDF, Word, or Excel",
        type=["pdf", "docx", "xlsx"],
    )

    if uploaded:
        file_name = safe_filename(uploaded.name)
        file_bytes = uploaded.getvalue()

else:
    st.info("For safety, the online URL mode accepts PDF URLs only.")
    url = st.text_input(
        "Direct PDF URL",
        placeholder="https://example.com/document.pdf",
    )

    if url and st.button("🌐 Download PDF"):
        try:
            file_bytes = download_document(url)
            file_name = safe_filename(
                os.path.basename(urlparse(url).path)
                or "downloaded_document.pdf"
            )
            st.success(f"Downloaded: {file_name}")
        except Exception as exc:
            st.error(str(exc))


# ============================================================
# ANALYZE
# ============================================================

if file_bytes is not None and file_name is not None:

    if len(file_bytes) > 25 * 1024 * 1024:
        st.error("Maximum file size is 25 MB.")
        st.stop()

    st.divider()
    st.subheader("📋 Selected Document")
    st.write(f"**File:** {file_name}")
    st.write(f"**Size:** {len(file_bytes) / 1024:.2f} KB")

    if st.button(
        "🔬 Analyze Document",
        type="primary",
        use_container_width=True,
    ):

        with st.spinner("Analyzing document safely..."):

            extension = os.path.splitext(file_name)[1].lower()

            try:
                documents, metadata = extract_content(
                    file_name,
                    file_bytes,
                )
            except Exception as exc:
                st.error(f"Could not parse document: {exc}")
                st.stop()

            file_hash = sha256_bytes(file_bytes)

            structure = []

            if do_structure:
                if extension == ".pdf":
                    structure = analyze_pdf_structure(file_bytes)
                else:
                    structure = analyze_office_structure(
                        file_bytes,
                        extension,
                    )

            iocs = extract_iocs(
                file_bytes,
                documents,
            ) if do_iocs else {
                "URLs": [],
                "IPv4": [],
                "SHA-256": [],
                "MD5": [],
            }

            clamav_result = (
                clamav_scan(file_bytes, file_name)
                if do_clamav
                else None
            )

            vt_result = (
                virustotal_hash_check(file_hash, virustotal_key)
                if do_vt
                else None
            )

            risk, score = risk_assessment(
                structure,
                clamav_result,
                vt_result,
            )

            st.session_state.analysis = {
                "file_name": file_name,
                "file_bytes": file_bytes,
                "documents": documents,
                "metadata": metadata,
                "hash": file_hash,
                "structure": structure,
                "iocs": iocs,
                "clamav": clamav_result,
                "virustotal": vt_result,
                "risk": risk,
                "score": score,
            }

        st.success("Analysis complete.")


# ============================================================
# RESULTS
# ============================================================

if "analysis" in st.session_state:

    result = st.session_state.analysis

    st.divider()
    st.header("📊 Investigation Results")

    c1, c2, c3, c4 = st.columns(4)

    with c1:
        st.metric("Risk", result["risk"])

    with c2:
        st.metric("Risk Score", result["score"])

    with c3:
        st.metric(
            "SHA-256",
            result["hash"][:12] + "…",
        )

    with c4:
        st.metric(
            "Extracted Chunks",
            len(result["documents"]),
        )

    if do_metadata:
        with st.expander("🧾 Metadata", expanded=True):
            st.json(result["metadata"])

    if do_hash:
        with st.expander("🔐 File Hash"):
            st.code(result["hash"])

    if do_structure:
        with st.expander("⚠️ Suspicious Structures", expanded=True):
            if result["structure"]:
                for item in result["structure"]:
                    st.warning(
                        f"**{item['indicator']}** — "
                        f"{item['description']} "
                        f"(count: {item['count']})"
                    )
            else:
                st.success(
                    "No commonly suspicious structures were detected."
                )

    if do_iocs:
        with st.expander("🎯 URLs / IOCs"):
            tabs = st.tabs(["URLs", "IPv4", "SHA-256", "MD5"])

            for tab, key in zip(
                tabs,
                ["URLs", "IPv4", "SHA-256", "MD5"],
            ):
                with tab:
                    values = result["iocs"][key]
                    if values:
                        for value in values:
                            st.code(value)
                    else:
                        st.info(f"No {key} indicators found.")

    if do_clamav:
        with st.expander("🦠 ClamAV"):
            scan = result["clamav"]
            if scan["status"] == "Clean":
                st.success("ClamAV did not detect a threat.")
            elif scan["status"] == "Threat Detected":
                st.error("ClamAV detected a potential threat.")
            else:
                st.warning(scan["status"])
            st.code(scan["details"])

    if do_vt:
        with st.expander("🌐 VirusTotal"):
            st.json(result["virustotal"])

    # --------------------------------------------------------
    # RAG
    # --------------------------------------------------------

    if agent_type in {
        "🤖 RAG Investigator",
        "🧩 Full Investigation",
    }:

        st.divider()
        st.header("🤖 RAG Investigator")

        model_name = st.selectbox(
            "Chat model",
            ["gpt-4o-mini", "gpt-4o"],
            key="rag_model",
        )

        if not openai_key:
            st.info(
                "Enter your OpenAI API key in the sidebar to use RAG."
            )
        elif not result["documents"]:
            st.warning(
                "No readable text was extracted, so RAG cannot answer "
                "content questions for this file."
            )
        else:

            if st.button(
                "🧠 Build RAG Knowledge Base",
                use_container_width=True,
            ):
                with st.spinner("Creating embeddings and FAISS index..."):
                    try:
                        st.session_state.vectorstore = build_vectorstore(
                            result["documents"],
                            openai_key,
                        )
                        st.session_state.chat_history = []
                        st.success("RAG knowledge base is ready.")
                    except Exception as exc:
                        st.error(f"Could not build RAG index: {exc}")

            if "vectorstore" in st.session_state:

                for msg in st.session_state.get(
                    "chat_history",
                    [],
                ):
                    with st.chat_message(msg["role"]):
                        st.markdown(msg["content"])

                question = st.chat_input(
                    "Ask a question about the document..."
                )

                if question:

                    st.session_state.chat_history.append(
                        {
                            "role": "user",
                            "content": question,
                        }
                    )

                    with st.chat_message("user"):
                        st.markdown(question)

                    with st.chat_message("assistant"):

                        with st.spinner("Investigating..."):

                            try:
                                answer, source_docs = answer_with_rag(
                                    st.session_state.vectorstore,
                                    question,
                                    openai_key,
                                    model_name,
                                )

                                st.markdown(answer)

                                with st.expander("🔎 Retrieved Evidence"):
                                    for doc in source_docs:
                                        st.markdown(
                                            f"**{doc.metadata}**"
                                        )
                                        st.text(
                                            doc.page_content[:700]
                                        )
                                        st.divider()

                                st.session_state.chat_history.append(
                                    {
                                        "role": "assistant",
                                        "content": answer,
                                    }
                                )

                            except Exception as exc:
                                st.error(
                                    f"RAG error: {exc}"
                                )

else:
    st.info(
        "👆 Upload a PDF, DOCX, or XLSX file to begin."
    )
